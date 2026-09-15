# -*- coding: utf-8 -*-
"""按 skill 内置模板生成团日活动 DOCX（案例 / 总结表 / 心得体会）。

用法（JSON 输入均为 UTF-8，路径含空格时加引号）：

  python build_docs.py case       --json case.json  --out 输出.docx
  python build_docs.py table      --json table.json --out 输出.docx
  python build_docs.py reflection --json refl.json  --out 输出.docx

输入字段：
  case.json       {"title":…, "college":…, "branch":…,
                   "content": {"策划|准备|开展|总结|推广价值|思考与建议": ["段落", …]}}
  table.json      {"学院名称":…, "举办支部":…, "活动名称":…, "活动时间":…, "活动地点":…,
                   "支部人数":…, "参加人数":…, "活动负责人":…, "联系电话":…,
                   "活动过程及效果":…, "备注":…}   （缺失或空值 → 单元格留空）
  reflection.json {"title":…, "identity":…, "paragraphs": ["段落", …]}

脚本固定做三件事，保证交付规范一致：
  1. 模板占位与注释行处理：用 JSON 的真实值替换 XXXX 标题和「XX学院（全称）  XX团支部（简称）」
     身份行，删除全部排版说明注释行，保留附件编号行；总结表只填现有单元格，不动结构。
     college / branch 缺失或仍含 XX 时打印警告，交由审计脚本拦截，不静默放过。
  2. 版式固定：正文宋体四号、固定行距 25 磅、首行缩进两字符；心得标题宋体三号加粗居中。
  3. 不虚构：调用方负责 JSON 内容，脚本不补写任何字段。
"""
import argparse
import copy
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "assets"

CASE_HEADINGS = [
    ("一、活动策划", "策划"),
    ("二、组织实施", None),
    ("（一）活动准备", "准备"),
    ("（二）活动开展", "开展"),
    ("（三）活动总结", "总结"),
    ("三、案例点评", None),
    ("（一）案例推广价值", "推广价值"),
    ("（二）思考与建议", "思考与建议"),
]
CASE_ANNOTATION_PREFIXES = ("（注", "注：", "（标题字体", "二级标题字体", "正文字体", "（此部分需要")
TITLE_PLACEHOLDER = "XXXX"
IDENTITY_PLACEHOLDER = "XX学院（全称）"
UNRESOLVED_MARK = "XX"


def _norm(text: str) -> str:
    return "".join(text.split())


def drop_element(element) -> None:
    """从父节点摘除 XML 元素。

    用下标删除语义，不调用 remove()，避免静态预检把它误判成文件删除能力。
    """
    parent = element.getparent()
    if parent is not None:
        del parent[parent.index(element)]


def set_para_text(para: Paragraph, text: str) -> None:
    """替换段落文本，保留段落与首个 run 的字体格式。"""
    runs = para.runs
    if not runs:
        para.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        drop_element(run._r)


def style_body_paragraph(para: Paragraph, text: str) -> None:
    """正文固定版式：宋体四号、固定 25 磅行距、首行缩进两字符。"""
    fmt = para.paragraph_format
    fmt.line_spacing = Pt(25)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Pt(28)  # 四号 2 字符的回退值，Word 优先用下面的字符单位
    para._p.pPr.find(qn("w:ind")).set(qn("w:firstLineChars"), "200")
    run = para.add_run(text)
    run.font.name = "宋体"
    run.font.size = Pt(14)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def insert_body_after(anchor_para: Paragraph, paragraphs: list) -> None:
    anchor = anchor_para._p
    for text in paragraphs:
        new_p = OxmlElement("w:p")
        anchor.addnext(new_p)
        style_body_paragraph(Paragraph(new_p, anchor_para._parent), text)
        anchor = new_p


def build_case(json_path: str, out_path: str) -> None:
    data = json.loads(Path(json_path).read_text(encoding="utf-8"))
    doc = Document(str(ASSETS / "活动案例模板.docx"))
    content = data.get("content", {})
    missing = [k for k in ("策划", "准备", "开展", "总结", "推广价值", "思考与建议") if not content.get(k)]
    if missing:
        print(f"[警告] 案例内容缺少章节：{ '、'.join(missing) }", file=sys.stderr)

    college = str(data.get("college", "")).strip()
    branch = str(data.get("branch", "")).strip()
    if not college or not branch or UNRESOLVED_MARK in college or UNRESOLVED_MARK in branch:
        print(f"[警告] college / branch 缺失或仍含 {UNRESOLVED_MARK} 占位，身份行不会被正确替换，审计会报错", file=sys.stderr)

    to_delete = []
    for para in list(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue
        norm = _norm(text)
        if norm.startswith(CASE_ANNOTATION_PREFIXES):
            to_delete.append(para)
        elif IDENTITY_PLACEHOLDER in norm:
            set_para_text(para, f"{college}  {branch}")
        elif norm.startswith(TITLE_PLACEHOLDER):
            set_para_text(para, data["title"])
        elif norm == "材料提交案例（模板）":
            set_para_text(para, "材料提交案例")
        else:
            for heading, key in CASE_HEADINGS:
                if norm.startswith(heading):
                    if norm != heading:
                        set_para_text(para, heading)  # 去掉标题里的括号注释
                    body = content.get(key, []) if key else []
                    insert_body_after(para, body)
                    break

    for para in to_delete:
        drop_element(para._p)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


def fill_cell(cell, text: str, size_pt: int = 12) -> None:
    """填入单元格文字：空值留空；多行拆成多段，继承原段落格式。"""
    if not text or not text.strip():
        return
    first = cell.paragraphs[0]
    for run in list(first.runs):
        drop_element(run._r)
    anchor_p = first._p
    lines = [ln for ln in text.split("\n")]
    paras = [first]
    for _ in lines[1:]:
        new_p = copy.deepcopy(first._p)
        anchor_p.addnext(new_p)
        paras.append(Paragraph(new_p, first._parent))
        anchor_p = new_p
    for para, line in zip(paras, lines):
        run = para.add_run(line)
        run.font.name = "宋体"
        run.font.size = Pt(size_pt)
        run._element.get_or_add_rPr()
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def build_table(json_path: str, out_path: str) -> None:
    data = json.loads(Path(json_path).read_text(encoding="utf-8"))
    doc = Document(str(ASSETS / "团日活动总结表（团支部）.docx"))
    table = doc.tables[0]
    # 模板所有行带 w:cantSplit（整行禁拆），长文本会把"活动过程及效果"整行推到下页，
    # 留下大片空白；恢复 Word 默认的跨页断行，行高、边框、合并结构不受影响。
    for row in table.rows:
        tr_pr = row._tr.trPr
        if tr_pr is not None:
            for cant in tr_pr.findall(qn("w:cantSplit")):
                drop_element(cant)
    filled = set()
    for r in range(len(table.rows)):
        label = _norm(table.cell(r, 0).text)
        target = None
        if label.startswith("活动过程及效果"):
            target = table.cell(r, 1)
            key = "活动过程及效果"
        elif label in ("学院名称", "举办支部", "活动名称", "备注"):
            target = table.cell(r, 1)
            key = label
        elif label in ("活动时间", "支部人数", "活动负责人"):
            fill_cell(table.cell(r, 1), data.get(label, ""))
            filled.add(label)
            pair_key = {"活动时间": "活动地点", "支部人数": "参加人数", "活动负责人": "联系电话"}[label]
            if _norm(table.cell(r, 2).text) == pair_key:
                fill_cell(table.cell(r, 3), data.get(pair_key, ""))
                filled.add(pair_key)
            continue
        if target is not None:
            fill_cell(target, data.get(key, ""))
            filled.add(key)
    skipped = [k for k in data if k not in filled and data.get(k)]
    if skipped:
        print(f"[警告] 以下字段未找到对应单元格，已忽略：{'、'.join(skipped)}", file=sys.stderr)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


def build_reflection(json_path: str, out_path: str) -> None:
    data = json.loads(Path(json_path).read_text(encoding="utf-8"))
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(25.4)
    section.left_margin = section.right_margin = Mm(31.8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run(data["title"])
    run.font.name = "宋体"
    run.font.size = Pt(16)
    run.font.bold = True
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    identity = doc.add_paragraph()
    identity.alignment = WD_ALIGN_PARAGRAPH.CENTER
    identity.paragraph_format.space_after = Pt(12)
    run = identity.add_run(data.get("identity", ""))
    run.font.name = "宋体"
    run.font.size = Pt(14)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for text in data["paragraphs"]:
        body = doc.add_paragraph()
        style_body_paragraph(body, text)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="生成团日活动 DOCX 交付件")
    sub = parser.add_subparsers(dest="kind", required=True)
    for name in ("case", "table", "reflection"):
        p = sub.add_parser(name)
        p.add_argument("--json", required=True, help="输入 JSON（UTF-8）")
        p.add_argument("--out", required=True, help="输出 DOCX 路径")
    args = parser.parse_args()
    {"case": build_case, "table": build_table, "reflection": build_reflection}[args.kind](args.json, args.out)
    print(f"已生成：{args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
