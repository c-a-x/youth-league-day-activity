# -*- coding: utf-8 -*-
"""对生成的团日活动 DOCX 做启发式版式审计（参照 csu-thesis-format 的检查脚本模式）。

用法：
  python check_docx_format.py 文件1.docx 文件2.docx [--kind case|table|reflection|auto]

检查内容：
  1. 禁留物扫描：[待补]、XXXX、（模板）、模板排版注释行、未替换的「XX学院（全称）」占位。
  2. case：必含章节标题及其相对顺序；正文段落版式（宋体四号、固定 25 磅、首行缩进两字符）。
  3. reflection：标题三号加粗居中、身份行四号居中、正文版式同上。
  4. table：字段标签齐全；学院团委意见栏不得填写内容；禁止占位符。
  5. 页面为 A4 纵向。
输出 ERROR/WARN/INFO 分级报告；存在 ERROR 时退出码为 1。渲染检查另由 render_preview.py 负责。
"""
import argparse
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

FORBIDDEN = ["[待补]", "XXXX", "（模板）", "注：一级标题", "（注：", "（标题字体", "（此部分需要", "XX学院（全称）"]
CASE_HEADINGS = ["一、活动策划", "二、组织实施", "（一）活动准备", "（二）活动开展", "（三）活动总结", "三、案例点评", "（一）案例推广价值", "（二）思考与建议"]
TABLE_LABELS = ["学院名称", "举办支部", "活动名称", "活动时间", "活动地点", "支部人数", "参加人数", "活动负责人", "联系电话", "备注"]

errors: list[str] = []
warns: list[str] = []
infos: list[str] = []


def err(msg): errors.append(msg)
def warn(msg): warns.append(msg)
def info(msg): infos.append(msg)


def norm(text: str) -> str:
    return "".join(text.split())


def detect_kind(doc) -> str:
    texts = [norm(p.text) for p in doc.paragraphs]
    if any("主题团日活动案例" in t for t in texts):
        return "case"
    if doc.tables and any(norm("活动过程及效果").startswith(norm(c.text)[:4]) for c in doc.tables[0].column_cells(0) if c.text):
        return "table"
    return "reflection"


def check_forbidden(all_text: str) -> None:
    for token in FORBIDDEN:
        if token in all_text:
            err(f"发现禁留占位/注释：{token}")


def check_run_font(para, expect_size_pt, label, expect_bold=None, check_indent=False):
    """校验段落正文 run 的字体/字号/行距/缩进。"""
    text = para.text.strip()
    if not text:
        return
    pf = para.paragraph_format
    if pf.line_spacing is None or pf.line_spacing_rule is None or pf.line_spacing != __import__("docx").shared.Pt(25):
        err(f"{label}「{text[:12]}…」行距不是固定 25 磅")
    if check_indent:
        ind = para._p.pPr.find(qn("w:ind")) if para._p.pPr is not None else None
        if ind is None or ind.get(qn("w:firstLineChars")) != "200":
            err(f"{label}「{text[:12]}…」首行未缩进两字符")
    for run in para.runs:
        east = run._element.rPr.rFonts.get(qn("w:eastAsia")) if run._element.rPr is not None and run._element.rPr.rFonts is not None else None
        if east != "宋体":
            err(f"{label}「{text[:12]}…」中文字体不是宋体（{east}）")
            break
    sizes = {run.font.size.pt for run in para.runs if run.font.size is not None}
    if sizes and sizes != {expect_size_pt}:
        err(f"{label}「{text[:12]}…」字号 {sizes} 应为 {expect_size_pt}pt")
    if expect_bold is not None:
        bolds = {bool(run.font.bold) for run in para.runs}
        if bolds and bool(para.runs[0].font.bold) != expect_bold:
            warn(f"{label}「{text[:12]}…」加粗状态 {bolds} 与预期 {expect_bold} 不符")


HEADING_PREFIXES = ("一、", "二、", "三、", "（一）", "（二）", "（三）")


def check_case(doc):
    positions = []
    texts = [norm(p.text) for p in doc.paragraphs]
    for heading in CASE_HEADINGS:
        if heading not in texts:
            err(f"缺少章节标题：{heading}")
        else:
            positions.append((texts.index(heading), heading))
    order = [h for _, h in sorted(positions)]
    expected = [h for h in CASE_HEADINGS if h in order]
    if order != expected:
        err(f"章节顺序异常：{order}")
    body_count = 0
    for para in doc.paragraphs:
        text = norm(para.text)
        if not text:
            continue
        if any(text.startswith(h) for h in CASE_HEADINGS) or text in ("附件3：", "材料提交案例"):
            continue
        if text.startswith("“国家安全") and text.endswith("案例"):
            continue  # 案例主标题
        if len(text) < 30 and text.endswith(("支部", "团支部")):
            continue  # 身份行
        body_count += 1
        check_run_font(para, 14, "案例正文", expect_bold=False, check_indent=True)
    info(f"案例正文段落 {body_count} 段")


def check_reflection(doc):
    paras = [p for p in doc.paragraphs if p.text.strip()]
    if not paras:
        err("文档为空")
        return
    title, identity = paras[0], paras[1] if len(paras) > 1 else None
    if not title.text.strip().endswith(("心得体会", "读后感")):
        warn(f"首段不像标题：{title.text.strip()[:20]}")
    from docx.shared import Pt
    for run in title.runs:
        if run.font.size is not None and run.font.size != Pt(16):
            err(f"心得标题字号 {run.font.size.pt} 应为 16pt（三号）")
        if not run.font.bold:
            err("心得标题未加粗")
    if title.alignment is None or "CENTER" not in str(title.alignment):
        err("心得标题未居中")
    if identity is not None:
        if identity.alignment is None or "CENTER" not in str(identity.alignment):
            err("身份行未居中")
        for run in identity.runs:
            if run.font.size is not None and run.font.size != Pt(14):
                err(f"身份行字号 {run.font.size.pt} 应为 14pt（四号）")
    for para in paras[2:]:
        check_run_font(para, 14, "心得正文", expect_bold=False, check_indent=True)


def check_table(doc):
    table = doc.tables[0]
    labels_seen = {norm(cell.text) for row in table.rows for cell in row.cells}
    for label in TABLE_LABELS:
        if not any(lb.startswith(label) for lb in labels_seen):
            err(f"缺少字段行：{label}")
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            text = cell.text
            for token in FORBIDDEN:
                if token in text:
                    err(f"表格({r},{c})发现禁留占位：{token}")
            if norm(cell.text).startswith("学院团委意见") or (r > 0 and "团委" in norm(table.cell(r - 1, 0).text)):
                continue
    # 学院团委意见栏：除“盖章/年 月 日”外不得有内容
    for r in range(len(table.rows)):
        if "学院团委意见" in norm(table.cell(r, 0).text):
            content = norm(table.cell(r, 1).text).replace("盖章", "").replace("年月日", "")
            if content.strip():
                err("学院团委意见栏被填写，应保留空白")
    info("总结表字段检查完成")


def check_page(doc, name):
    sec = doc.sections[0]
    w_mm = sec.page_width.mm if sec.page_width else 0
    h_mm = sec.page_height.mm if sec.page_height else 0
    if not (205 <= w_mm <= 215 and 292 <= h_mm <= 302):
        err(f"{name} 页面不是 A4 纵向（{w_mm:.0f}x{h_mm:.0f}mm）")


def main() -> int:
    parser = argparse.ArgumentParser(description="团日活动 DOCX 版式审计")
    parser.add_argument("docx", nargs="+")
    parser.add_argument("--kind", choices=["case", "table", "reflection", "auto"], default="auto")
    args = parser.parse_args()

    for path in map(Path, args.docx):
        name = path.name
        doc = Document(str(path))
        kind = args.kind if args.kind != "auto" else detect_kind(doc)
        print(f"== {name}（识别为 {kind}）")
        all_text = "\n".join(p.text for p in doc.paragraphs)
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    all_text += "\n" + cell.text
        check_forbidden(all_text)
        check_page(doc, name)
        if kind == "case":
            check_case(doc)
        elif kind == "reflection":
            check_reflection(doc)
        elif kind == "table":
            check_table(doc)

    for msg in errors:
        print(f"[ERROR] {msg}")
    for msg in warns:
        print(f"[WARN] {msg}")
    for msg in infos:
        print(f"[INFO] {msg}")
    print(f"结果：{len(errors)} 个错误，{len(warns)} 个警告")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
